"""Construtor de apresentações PowerPoint (PPTX) com python-pptx e cálculos geométricos."""

import io
from pathlib import Path
from typing import Literal

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from config import presets
from core.engines.image import load_image_with_orientation_and_rotation
from core.engines.pdf import render_pdf_page_to_image
from core.models.domain import SlideGeometry, SlideItem


def parse_hex_color(hex_str: str) -> RGBColor:
    """Converte string hexadecimal (#RRGGBB ou RRGGBB) para objeto RGBColor do python-pptx."""
    cleaned = hex_str.lstrip("#")
    if len(cleaned) == 3:
        cleaned = "".join([c * 2 for c in cleaned])
    if len(cleaned) != 6:
        return RGBColor(255, 255, 255)
    r = int(cleaned[0:2], 16)
    g = int(cleaned[2:4], 16)
    b = int(cleaned[4:6], 16)
    return RGBColor(r, g, b)


def calculate_slide_geometry(
    img_width: int,
    img_height: int,
    slide_width_inches: float,
    slide_height_inches: float,
    fit_mode: Literal["fit", "fill", "stretch"] = "fit",
) -> SlideGeometry:
    """Calcula as coordenadas de posicionamento e dimensões da imagem no slide."""
    if img_width <= 0 or img_height <= 0:
        return SlideGeometry(
            slide_width_inches=slide_width_inches,
            slide_height_inches=slide_height_inches,
            image_width=img_width,
            image_height=img_height,
            left_inches=0.0,
            top_inches=0.0,
            render_width_inches=slide_width_inches,
            render_height_inches=slide_height_inches,
        )

    img_aspect = img_width / img_height
    slide_aspect = slide_width_inches / slide_height_inches

    if fit_mode == "stretch":
        return SlideGeometry(
            slide_width_inches=slide_width_inches,
            slide_height_inches=slide_height_inches,
            image_width=img_width,
            image_height=img_height,
            left_inches=0.0,
            top_inches=0.0,
            render_width_inches=slide_width_inches,
            render_height_inches=slide_height_inches,
        )

    if fit_mode == "fill":
        # Cobre todo o slide (imagem pode sangrar para fora das margens)
        if img_aspect > slide_aspect:
            render_h = slide_height_inches
            render_w = render_h * img_aspect
            top = 0.0
            left = (slide_width_inches - render_w) / 2.0
        else:
            render_w = slide_width_inches
            render_h = render_w / img_aspect
            left = 0.0
            top = (slide_height_inches - render_h) / 2.0

        return SlideGeometry(
            slide_width_inches=slide_width_inches,
            slide_height_inches=slide_height_inches,
            image_width=img_width,
            image_height=img_height,
            left_inches=left,
            top_inches=top,
            render_width_inches=render_w,
            render_height_inches=render_h,
        )

    # Modo padrão: "fit" (mantém proporção e centraliza sem cortes)
    if img_aspect > slide_aspect:
        render_w = slide_width_inches
        render_h = render_w / img_aspect
        left = 0.0
        top = (slide_height_inches - render_h) / 2.0
    else:
        render_h = slide_height_inches
        render_w = render_h * img_aspect
        top = 0.0
        left = (slide_width_inches - render_w) / 2.0

    return SlideGeometry(
        slide_width_inches=slide_width_inches,
        slide_height_inches=slide_height_inches,
        image_width=img_width,
        image_height=img_height,
        left_inches=left,
        top_inches=top,
        render_width_inches=render_w,
        render_height_inches=render_h,
    )


def build_presentation(
    items: list[SlideItem],
    output_path: Path | str,
    aspect_ratio: str = "16:9",
    fit_mode: Literal["fit", "fill", "stretch"] = "fit",
    bg_color: str = "#FFFFFF",
    pdf_dpi: int = 200,
) -> Path:
    """Gera uma apresentação PowerPoint completa contendo os itens fornecidos."""
    output_file = Path(output_path)
    if not output_file.name.lower().endswith(".pptx"):
        output_file = output_file.with_suffix(".pptx")

    output_file.parent.mkdir(parents=True, exist_ok=True)

    aspect_cfg = presets.aspect_ratios.get(aspect_ratio, presets.aspect_ratios["16:9"])
    slide_w_in = float(aspect_cfg.width_inches)
    slide_h_in = float(aspect_cfg.height_inches)

    prs = Presentation()
    prs.slide_width = Inches(slide_w_in)
    prs.slide_height = Inches(slide_h_in)

    blank_slide_layout = prs.slide_layouts[6]
    is_transparent = str(bg_color).strip().lower() in ("transparent", "none", "")
    background_rgb = None if is_transparent else parse_hex_color(bg_color)

    active_items = [it for it in items if it.enabled]

    for item in active_items:
        slide = prs.slides.add_slide(blank_slide_layout)

        if not is_transparent and background_rgb is not None:
            bg = slide.background
            fill = bg.fill
            fill.solid()
            fill.fore_color.rgb = background_rgb

        img_buffer = io.BytesIO()
        if item.source_type == "pdf_page":
            page_idx = item.page_index if item.page_index is not None else 0
            pil_img = render_pdf_page_to_image(
                item.source_path,
                page_index=page_idx,
                dpi=pdf_dpi,
                rotation=item.rotation,
            )
            pil_img.save(img_buffer, format="PNG", optimize=False)
            current_w, current_h = pil_img.width, pil_img.height
        else:
            pil_img = load_image_with_orientation_and_rotation(
                item.source_path, rotation=item.rotation
            )
            if pil_img.mode in ("RGBA", "LA", "P"):
                pil_img.save(img_buffer, format="PNG")
            else:
                pil_img.save(img_buffer, format="JPEG", quality=95)
            current_w, current_h = pil_img.width, pil_img.height

        img_buffer.seek(0)

        geom = calculate_slide_geometry(
            img_width=current_w,
            img_height=current_h,
            slide_width_inches=slide_w_in,
            slide_height_inches=slide_h_in,
            fit_mode=fit_mode,
        )

        slide.shapes.add_picture(
            img_buffer,
            left=Inches(geom.left_inches),
            top=Inches(geom.top_inches),
            width=Inches(geom.render_width_inches),
            height=Inches(geom.render_height_inches),
        )

    prs.save(str(output_file))
    return output_file
