"""Testes para o construtor de apresentações PowerPoint (PPTX)."""

from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from core.engines.pptx import (
    build_presentation,
    calculate_slide_geometry,
    parse_hex_color,
)
from core.models import SlideItem


def test_parse_hex_color() -> None:
    """Verifica conversão de strings hex para RGBColor com casos normais e especiais."""
    c1 = parse_hex_color("#FFFFFF")
    assert (c1[0], c1[1], c1[2]) == (255, 255, 255)

    c2 = parse_hex_color("#000000")
    assert (c2[0], c2[1], c2[2]) == (0, 0, 0)

    c3 = parse_hex_color("6366F1")
    assert (c3[0], c3[1], c3[2]) == (0x63, 0x66, 0xF1)

    # 3 dígitos (#FFF -> #FFFFFF)
    c4 = parse_hex_color("#ABC")
    assert (c4[0], c4[1], c4[2]) == (0xAA, 0xBB, 0xCC)

    # String inválida retorna branco padrão
    c_inv = parse_hex_color("invalido")
    assert (c_inv[0], c_inv[1], c_inv[2]) == (255, 255, 255)


def test_calculate_slide_geometry_modes() -> None:
    """Valida o cálculo de geometria em todos os modos: fit, fill, stretch e dimensões zeradas."""
    # 1. Dimensão zero / inválida
    geom_zero = calculate_slide_geometry(0, 0, 13.333, 7.5, "fit")
    assert geom_zero.render_width_inches == 13.333
    assert geom_zero.render_height_inches == 7.5

    # 2. Stretch
    geom_str = calculate_slide_geometry(500, 1000, 13.333, 7.5, "stretch")
    assert geom_str.render_width_inches == 13.333
    assert geom_str.render_height_inches == 7.5

    # 3. Fill mais largo que o slide (img_aspect > slide_aspect)
    geom_fill_w = calculate_slide_geometry(3000, 1000, 13.333, 7.5, "fill")
    assert pytest.approx(geom_fill_w.render_height_inches, abs=0.01) == 7.5
    assert geom_fill_w.render_width_inches > 13.333

    # 4. Fill mais alto que o slide (img_aspect <= slide_aspect)
    geom_fill_h = calculate_slide_geometry(1000, 2000, 13.333, 7.5, "fill")
    assert pytest.approx(geom_fill_h.render_width_inches, abs=0.01) == 13.333
    assert geom_fill_h.render_height_inches > 7.5

    # 5. Fit padrão (16:9 em 16:9)
    geom = calculate_slide_geometry(1920, 1080, 13.333, 7.5, "fit")
    assert pytest.approx(geom.left_inches, abs=0.01) == 0.0
    assert pytest.approx(geom.top_inches, abs=0.01) == 0.0
    assert pytest.approx(geom.render_width_inches, abs=0.01) == 13.333
    assert pytest.approx(geom.render_height_inches, abs=0.01) == 7.5

    # 6. Fit quadrado (limitado por altura)
    geom_sq = calculate_slide_geometry(1000, 1000, 13.333, 7.5, "fit")
    assert pytest.approx(geom_sq.render_height_inches, abs=0.01) == 7.5
    assert pytest.approx(geom_sq.render_width_inches, abs=0.01) == 7.5
    assert geom_sq.left_inches > 0.0


def test_build_presentation_from_images(tmp_path: Path) -> None:
    """Cria uma apresentação PPTX a partir de imagens e verifica a integridade do arquivo gerado."""
    items: list[SlideItem] = []
    for i in range(1, 4):
        img_path = tmp_path / f"slide_{i}.png"
        img = Image.new("RGB", (1280, 720), color=(i * 60, 100, 150))
        img.save(img_path)

        item = SlideItem(
            id=f"item_{i}",
            source_type="image",
            source_path=str(img_path),
            title=f"Slide {i}",
            original_width=1280,
            original_height=720,
        )
        items.append(item)

    # Item desabilitado para testar filtro de enabled
    items.append(
        SlideItem(
            id="item_disabled",
            source_type="image",
            source_path=str(tmp_path / "slide_1.png"),
            enabled=False,
        )
    )

    # Caminho sem extensão .pptx para testar adição automática
    out_file = tmp_path / "Teste_Apresentacao"
    saved_path = build_presentation(
        items=items,
        output_path=out_file,
        aspect_ratio="16:9",
        fit_mode="fit",
        bg_color="#000000",
    )

    assert saved_path.exists()
    assert str(saved_path).endswith(".pptx")
    assert saved_path.stat().st_size > 1000

    prs = Presentation(str(saved_path))
    assert len(prs.slides) == 3
    assert prs.slide_width is not None and prs.slide_height is not None
    assert pytest.approx(prs.slide_width.inches, abs=0.01) == 13.333
    assert pytest.approx(prs.slide_height.inches, abs=0.01) == 7.5


def test_build_presentation_from_pdf(sample_pdf: Path, tmp_path: Path) -> None:
    """Gera uma apresentação PPTX a partir de páginas de um arquivo PDF."""
    items = [
        SlideItem(
            id=f"pdf_page_{idx}",
            source_type="pdf_page",
            source_path=str(sample_pdf.resolve()),
            page_index=idx,
            title=f"Página {idx + 1}",
        )
        for idx in range(2)
    ]

    out_file = tmp_path / "PDF_Apresentacao.pptx"
    saved = build_presentation(items, out_file, aspect_ratio="16:9", pdf_dpi=150)

    assert saved.exists()
    prs = Presentation(str(saved))
    assert len(prs.slides) == 2


def test_build_presentation_transparent_background(tmp_path: Path) -> None:
    """Verifica que a opção bg_color='transparent' cria apresentação válida sem erro."""
    img_path = tmp_path / "slide_transp.png"
    img = Image.new("RGBA", (800, 600), color=(100, 200, 250, 200))
    img.save(img_path)

    item = SlideItem(
        id="item_transp",
        source_type="image",
        source_path=str(img_path),
        title="Slide Transparente",
        original_width=800,
        original_height=600,
    )

    out_file = tmp_path / "Apresentacao_Transparente.pptx"
    saved = build_presentation(
        items=[item],
        output_path=out_file,
        aspect_ratio="16:9",
        bg_color="transparent",
    )

    assert saved.exists()
    prs = Presentation(str(saved))
    assert len(prs.slides) == 1
